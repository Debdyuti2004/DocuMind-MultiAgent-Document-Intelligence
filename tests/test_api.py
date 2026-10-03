"""API endpoint tests using FastAPI TestClient."""

from pathlib import Path
from io import BytesIO
import pytest
import docx
import pandas as pd
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)
SAMPLE_DIR = Path("data/sample_docs")


@pytest.fixture
def uploaded_test_documents():
    """Remove test uploads so API tests leave the user's index unchanged."""
    document_ids = []
    yield document_ids
    for document_id in document_ids:
        response = client.delete(f"/documents/{document_id}")
        assert response.status_code == 200, response.text


def test_api_root_and_health():
    r1 = client.get("/")
    assert r1.status_code == 200
    assert r1.json()["system"] == "DocuMind — Multi-Agent Document Intelligence"

    r2 = client.get("/health")
    assert r2.status_code == 200
    assert r2.json()["status"] == "healthy"


def test_api_upload_and_documents(uploaded_test_documents):
    pdf_path = SAMPLE_DIR / "ApexCorp_Annual_Report_FY2024.pdf"
    with open(pdf_path, "rb") as f:
        resp = client.post("/upload", files={"file": (pdf_path.name, f, "application/pdf")})
    
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "success"
    doc_id = data["metadata"]["document_id"]
    uploaded_test_documents.append(doc_id)
    assert doc_id.startswith("doc_")
    assert data["summary"]["document_type"] == "Annual Report"

    # Verify listing
    list_resp = client.get("/documents")
    assert list_resp.status_code == 200
    docs = list_resp.json()
    assert any(d["document_id"] == doc_id for d in docs)


def test_api_upload_supported_text_and_spreadsheet_formats(uploaded_test_documents):
    word_buffer = BytesIO()
    word = docx.Document()
    word.add_heading("Upload verification", level=1)
    word.add_paragraph("The document states a revenue figure of INR 42 Crore.")
    word.save(word_buffer)

    workbook_buffer = BytesIO()
    pd.DataFrame({"Metric": ["Revenue"], "Value": [42]}).to_excel(workbook_buffer, index=False)

    examples = [
        ("upload-check.txt", "text/plain", b"Upload verification. Revenue is INR 42 Crore."),
        ("upload-check.csv", "text/csv", b"Metric,Value\nRevenue,42\n"),
        ("upload-check.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", word_buffer.getvalue()),
        ("upload-check.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", workbook_buffer.getvalue()),
    ]
    for filename, content_type, content in examples:
        response = client.post("/upload", files={"file": (filename, content, content_type)})
        assert response.status_code == 200, f"{filename}: {response.text}"
        data = response.json()
        uploaded_test_documents.append(data["metadata"]["document_id"])
        assert data["metadata"]["num_chunks"] > 0
        if filename.endswith((".csv", ".xlsx")):
            assert data["metadata"]["num_tables"] > 0


def test_api_rejects_unsupported_empty_and_blank_requests():
    unsupported = client.post("/upload", files={"file": ("not-a-document.exe", b"no", "application/octet-stream")})
    assert unsupported.status_code == 415

    empty_file = client.post("/upload", files={"file": ("empty.txt", b"", "text/plain")})
    assert empty_file.status_code == 400

    blank_query = client.post("/query", json={"query": "   "})
    assert blank_query.status_code == 422


def test_api_query():
    # Query across documents
    payload = {
        "query": "What was the Total Revenue in FY2024?",
        "enable_verification": True
    }
    resp = client.post("/query", json=payload)
    assert resp.status_code == 200
    data = resp.json()
    assert "answer" in data
    assert data["intent"] in ("financial_analysis", "document_qa")
    assert len(data["plan"]) > 0
    assert len(data["agent_logs"]) > 0
    assert data["confidence"] in ("High", "Medium", "Low")
    assert not any("change across selected documents" in item.get("metric", "") for item in data["derived_calculations"])
